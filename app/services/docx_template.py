"""Savollar uchun namuna Word fayl.

Admin «Docx yuklash» ni bosganda shu fayl yuboriladi — format xatolari
(ball, vaqt, rasm, Lehine/Aleyhine jadvali) kamayadi. Fayl kod bilan
yaratiladi, shuning uchun har doim parser bilan mos keladi
(tests/test_docx_template.py tekshiradi).
"""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Pt

TEMPLATE_FILENAME = "savollar_namuna.docx"


def _placeholder_image() -> BytesIO:
    """Rasm o'rnini ko'rsatuvchi oddiy PNG (haqiqiy rasm bilan almashtiriladi)."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (480, 240), (229, 237, 247))
    draw = ImageDraw.Draw(img)
    draw.rectangle([4, 4, 475, 235], outline=(37, 99, 235), width=4)
    draw.text((150, 110), "Rasm shu yerga qo'yiladi", fill=(30, 41, 59))
    buf = BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def _question(doc, section: str, text: str, prep: int, answer: int, points: int,
              numbered: bool = True, image: bool = False) -> None:
    doc.add_paragraph(f"Bölüm {section}")
    doc.add_paragraph(f"1.soru: {text}" if numbered else text)
    doc.add_paragraph(f"Tayyorlanish: {prep}")
    doc.add_paragraph(f"Javob: {answer}")
    doc.add_paragraph(f"Bal: {points}")
    if image:
        doc.add_picture(_placeholder_image())


def build_questions_template(path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)

    doc.add_heading("Savollar namunasi", level=1)
    rules = doc.add_paragraph()
    rules.add_run("Qoidalar (bu qismni o'chirib, o'z savollaringizni yozing):\n").bold = True
    rules.add_run(
        "• Har bir bo'lim «Bölüm 1.1», «Bölüm 1.2» ... bilan boshlanadi.\n"
        "• Savol «1.soru:» bilan yoziladi (uzun bo'limlarda raqamsiz ham bo'ladi).\n"
        "• «Tayyorlanish: N» va «Javob: N» — soniyalarda.\n"
        "• «Bal: N» — savolning maksimal balli. Hamma ballar yig'indisi 75 bo'lsin.\n"
        "• Rasm savol tagiga qo'yiladi.\n"
        "• Munozara savolida 2 ustunli «Lehine | Aleyhine» jadvali bo'ladi."
    )

    _question(doc, "1.1", "Yaz tatilinde nereye gittiniz?", 5, 30, 10)
    _question(doc, "1.2", "Boş zamanlarınızda ne yapmayı seversiniz?", 5, 30, 12)
    _question(doc, "1.3", "Fotoğrafta neler görüyorsunuz?", 10, 45, 13, image=True)
    _question(
        doc, "2",
        "Bana bir şeye geç kaldığınız bir zamanı anlatın. Geç kalmanıza ne sebep oldu? "
        "Ne hissettiniz?",
        60, 120, 19, numbered=False, image=True,
    )
    _question(
        doc, "3",
        "Üniversite öğrencilerinin çalışması gerekir mi? Lehine ve aleyhine "
        "düşüncelerinizi anlatın.",
        60, 120, 21, numbered=False,
    )
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Lehine", "Aleyhine"),
        ("Maddi bağımsızlık kazandırır.", "Derslere zaman ayırmak zorlaşır."),
        ("İş tecrübesi kazanılır.", "Sağlığı olumsuz etkileyebilir."),
    ]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right

    doc.save(str(path))
    return path
