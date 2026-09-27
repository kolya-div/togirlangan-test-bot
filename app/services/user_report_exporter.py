import json
import logging
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models import Answer, Question, TestAttempt, User
from app.utils.helpers import format_filename, local_now, to_local

logger = logging.getLogger(__name__)

# Eksport qilingan hisobot fayllarining saqlanish papkasi (proyekt ildizi).
REPORTS_DIR = Path(__file__).resolve().parents[2] / "data" / "reports"


async def collect_users_with_results(
    session: AsyncSession,
) -> list[dict[str, Any]]:
    """BARCHA foydalanuvchilar (adminlar ham) va ularning barcha test
    attemptlarini to'liq ma'lumot bilan yig'adi.

    Wipe (to'liq tozalash) oldidan YAGONA ma'lumot manbai — .docx hisobot.
    Shuning uchun hech narsa filtirlanmaydi: har bir attempt alohida qator
    bo'ladi, test topshirmagan userlar ham "test topshirmagan" holatida
    kiradi (ularning to'liq ma'lumoti yo'qolmasligi uchun).
    """
    result = await session.execute(select(User).order_by(User.id))
    users = list(result.scalars().all())

    rows: list[dict[str, Any]] = []
    if not users:
        return rows

    user_ids = [u.id for u in users]
    att_result = await session.execute(
        select(TestAttempt).where(TestAttempt.user_id.in_(user_ids)).order_by(TestAttempt.id)
    )
    attempts = list(att_result.scalars().all())

    answers_by_attempt = await _collect_answers(session, [a.id for a in attempts])

    attempts_by_user: dict[Any, list[TestAttempt]] = {}
    for attempt in attempts:
        attempts_by_user.setdefault(attempt.user_id, []).append(attempt)

    for user in users:
        user_attempts = attempts_by_user.get(user.id, [])
        if not user_attempts:
            rows.append(
                {
                    "telegram_id": user.telegram_id,
                    "full_name": user.full_name or "*",
                    "username": user.username or "*",
                    "phone": user.phone or "*",
                    "created_at": user.created_at,
                    "status": "test topshirmagan",
                    "started_at": None,
                    "finished_at": None,
                    "score": None,
                    "level": None,
                    "answers": [],
                }
            )
            continue
        for attempt in user_attempts:
            rows.append(
                {
                    "telegram_id": user.telegram_id,
                    "full_name": user.full_name or "*",
                    "username": user.username or "*",
                    "phone": user.phone or "*",
                    "created_at": user.created_at,
                    "status": attempt.status,
                    "started_at": attempt.started_at,
                    "finished_at": attempt.finished_at,
                    "score": attempt.score,
                    "level": attempt.level,
                    "answers": answers_by_attempt.get(attempt.id, []),
                }
            )

    # Ism bo'yicha tartiblash (naklonistik, lekin amaliy jihatdan yetarli).
    rows.sort(key=lambda r: (r["full_name"] or "").lower())
    return rows


async def _collect_answers(
    session: AsyncSession,
    attempt_ids: list[int],
) -> dict[int, list[dict[str, Any]]]:
    """Har bir attempt uchun javoblar: savol, transkript, ball, xatolar.
    Savol tartibida (bo'lim, raqam) qaytariladi."""
    if not attempt_ids:
        return {}

    result = await session.execute(
        select(Answer, Question)
        .join(Question, Answer.question_id == Question.id, isouter=True)
        .where(Answer.attempt_id.in_(attempt_ids))
        .order_by(Answer.attempt_id, Question.section, Question.order_number)
    )

    by_attempt: dict[int, list[dict[str, Any]]] = {}
    for answer, question in result.all():
        feedback: dict[str, Any] = {}
        if answer.feedback:
            try:
                feedback = json.loads(answer.feedback)
            except (json.JSONDecodeError, TypeError):
                feedback = {}

        max_points = question.max_points if question else None
        earned = None
        if max_points and answer.score is not None:
            # report_service bilan bir xil hisob: foiz → savol balli
            earned = round(answer.score * max_points / 100)

        by_attempt.setdefault(answer.attempt_id, []).append(
            {
                "section": question.section if question else "",
                "order_number": question.order_number if question else 0,
                "question": question.text if question else "(savol o'chirilgan)",
                "transcript": answer.transcript or feedback.get("transcript") or "",
                "corrected_text": feedback.get("corrected_text") or "",
                "mistakes": feedback.get("mistakes") or [],
                "score": answer.score,
                "max_points": max_points,
                "earned": earned,
            }
        )
    return by_attempt


# XML (Word) da ruxsat etilmagan boshqaruv belgilari: \t, \n, \r dan tashqari.
# Bitta shunday belgi (ism yoki AI transkriptida) python-docx'da ValueError
# beradi va butun kunlik hisobot yuborilmay qoladi.
_XML_INVALID_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _clean(value: Any) -> str:
    """Matnni Word uchun xavfsiz qiladi."""
    if value is None:
        return ""
    return _XML_INVALID_RE.sub("", str(value))


def _fmt_dt(value) -> str:
    local = to_local(value)
    return local.strftime("%Y-%m-%d %H:%M") if local else "*"


def _fmt_answer_score(item: dict[str, Any]) -> str:
    if item["score"] is None:
        return "baholanmagan"
    if item["max_points"]:
        return f"{item['earned']}/{item['max_points']}"
    return f"{item['score']}%"


# ───────────────────────── Umumiy (Word + Excel) ─────────────────────────

MAX_SCORE = 75

GREEN = "166534"        # sarlavha / aksent
GREEN_LIGHT = "DCFCE7"  # savol sarlavhasi foni
ZEBRA = "F3F4F6"        # jadval qatorlari almashib
GRAY_TEXT = "6B7280"

LEVEL_COLORS = {        # daraja: (fon, matn)
    "C1": ("BBF7D0", "14532D"),
    "B2": ("BFDBFE", "1E3A8A"),
    "B1": ("FEF08A", "713F12"),
    "Below B1": ("FECACA", "7F1D1D"),
}

STATUS_LABELS = {
    "finished": "Yakunlangan",
    "processing": "Tekshirilmoqda",
    "active": "Tugatilmagan",
    "started": "Tugatilmagan",
    "test topshirmagan": "Test topshirmagan",
}


def _status_label(status: str | None) -> str:
    return STATUS_LABELS.get(status or "", status or "—")


def _sorted_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Ball bo'yicha kamayish tartibida; balli yo'qlar oxirida (ism bo'yicha)."""
    return sorted(
        rows,
        key=lambda r: (
            r["score"] is None,
            -(r["score"] or 0),
            (r["full_name"] or "").lower(),
        ),
    )


def _stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    users = {r["telegram_id"] for r in rows}
    scored = [r for r in rows if r["score"] is not None]
    levels = {lvl: sum(1 for r in scored if r["level"] == lvl) for lvl in LEVEL_COLORS}
    return {
        "users": len(users),
        "finished": sum(1 for r in rows if r["status"] == "finished"),
        "avg": round(sum(r["score"] for r in scored) / len(scored), 1) if scored else None,
        "levels": levels,
    }


def _question_columns(rows: list[dict[str, Any]]) -> list[tuple[str, int, Any]]:
    """Barcha javoblardagi savollar (bo'lim, tartib, max ball) — tartiblangan."""
    seen: dict[tuple[str, int], Any] = {}
    for r in rows:
        for a in r["answers"]:
            seen.setdefault((a["section"], a["order_number"]), a["max_points"])

    def key(k):
        parts = []
        for p in str(k[0]).split("."):
            parts.append(int(p) if p.isdigit() else 0)
        return (parts, k[1])

    return [(s, o, seen[(s, o)]) for (s, o) in sorted(seen, key=key)]


def _question_label(section: str, order: int, total_in_section: int) -> str:
    return f"{section}" if total_in_section == 1 else f"{section}.{order}"


def _labeler(rows: list[dict[str, Any]]):
    """(bo'lim, tartib) -> "1.4" (bo'limda bitta savol) yoki "2.1" (bir nechta)."""
    per_section: dict[str, int] = {}
    for s, _, _ in _question_columns(rows):
        per_section[s] = per_section.get(s, 0) + 1
    return lambda s, o: _question_label(s, o, per_section.get(s, 1))


# ───────────────────────────── Word ─────────────────────────────

def _shade(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    _insert_ordered(cell._tc.get_or_add_tcPr(), shd, _TC_PR_ORDER)


def _cell_text(cell, text: str, *, bold=False, color: str | None = None,
               size: float | None = None, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run(_clean(text))
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _set_widths(table, widths_cm: list[float]) -> None:
    """Ustun kengliklari: jadval to'ri (gridCol) ham, kataklar ham — faqat
    kataklar berilsa LibreOffice/Word to'rni teng bo'lib qo'yadi."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    table.autofit = False
    for col, w in zip(table.columns, widths_cm):
        col.width = Cm(w)
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)
    # Jadvalning umumiy kengligi va qat'iy (fixed) joylashuv. tblPr bolalari
    # sxemadagi tartibda bo'lishi shart — aks holda Word faylni "buzilgan"
    # deb ochishi mumkin.
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) / 2.54 * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    for el in (tbl_w, layout):
        _insert_ordered(tbl_pr, el, _TBL_PR_ORDER)


_TC_PR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]

_TBL_PR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
    "tblCaption", "tblDescription",
]


def _insert_ordered(parent, el, order: list[str]) -> None:
    """`el` ni (shu nomdagi eskisini almashtirib) sxema tartibida joylaydi."""
    from docx.oxml.ns import qn

    name = el.tag.split("}")[1]
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    later = set(order[order.index(name) + 1:])
    for i, child in enumerate(parent):
        if child.tag.split("}")[1] in later:
            parent.insert(i, el)
            return
    parent.append(el)


def _setup_document():
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for name, size in (("Heading 1", 18), ("Heading 2", 13)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(GREEN)
    return doc


def _add_stats_cards(doc, stats: dict[str, Any]) -> None:
    levels = " · ".join(f"{lvl}: {n}" for lvl, n in stats["levels"].items())
    cards = [
        ("Ishtirokchilar", str(stats["users"])),
        ("Yakunlagan", str(stats["finished"])),
        ("O'rtacha ball", f"{stats['avg']}/{MAX_SCORE}" if stats["avg"] is not None else "—"),
        ("Darajalar", levels),
    ]
    table = doc.add_table(rows=2, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(cards):
        top, bottom = table.rows[0].cells[i], table.rows[1].cells[i]
        _shade(top, GREEN_LIGHT)
        _shade(bottom, GREEN_LIGHT)
        _cell_text(top, label, color=GRAY_TEXT, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell_text(bottom, value, bold=True, color=GREEN,
                   size=14 if i < 3 else 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_widths(table, [5.5, 5.5, 5.5, 10.2])
    doc.add_paragraph()


def _add_summary_table(doc, rows: list[dict[str, Any]]) -> None:
    headers = ["#", "Ism familiya", "Telegram ID", "Telefon", "Username",
               "Holat", "Yakunlangan", "Ball", "Daraja"]
    widths = [1.0, 5.2, 3.0, 3.4, 3.6, 3.2, 3.2, 1.6, 2.5]  # jami 26.7 sm (A4 albom)
    center = {0, 2, 6, 7, 8}

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        _shade(cell, GREEN)
        _cell_text(cell, text, bold=True, color="FFFFFF", size=10,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    _repeat_header(table.rows[0])

    for idx, row in enumerate(rows, start=1):
        values = [
            str(idx),
            row["full_name"] or "—",
            str(row["telegram_id"]),
            row["phone"] or "—",
            f"@{row['username']}" if row["username"] and row["username"] != "*" else "—",
            _status_label(row["status"]),
            _fmt_dt(row["finished_at"]) if row["finished_at"] else "—",
            f"{row['score']}/{MAX_SCORE}" if row["score"] is not None else "—",
            row["level"] or "—",
        ]
        cells = table.add_row().cells
        for i, value in enumerate(values):
            _cell_text(cells[i], value, bold=(i in (1, 7)),
                       align=WD_ALIGN_PARAGRAPH.CENTER if i in center else None)
            if idx % 2 == 0:
                _shade(cells[i], ZEBRA)
        level_colors = LEVEL_COLORS.get(row["level"] or "")
        if level_colors:
            _shade(cells[8], level_colors[0])
            _cell_text(cells[8], row["level"], bold=True, color=level_colors[1],
                       align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_widths(table, widths)


def _add_answers_section(doc, rows: list[dict[str, Any]]) -> None:
    """Har bir foydalanuvchi alohida sahifada: har savol — ramkali blok."""
    with_answers = [r for r in rows if r["answers"]]
    if not with_answers:
        return
    label = _labeler(rows)

    for n, row in enumerate(with_answers, start=1):
        doc.add_page_break()
        score = f"{row['score']}/{MAX_SCORE}" if row["score"] is not None else "baholanmagan"
        doc.add_heading(_clean(f"{n}. {row['full_name']} — {score} · {row['level'] or '—'}"), level=2)
        meta = doc.add_paragraph()
        username = f"@{row['username']}" if row["username"] and row["username"] != "*" else "—"
        mrun = meta.add_run(_clean(
            f"Telegram ID: {row['telegram_id']}   ·   Telefon: {row['phone'] or '—'}   ·   "
            f"Username: {username}   ·   Yakunlangan: {_fmt_dt(row['finished_at'])}"
        ))
        mrun.font.size = Pt(9)
        mrun.font.color.rgb = RGBColor.from_string(GRAY_TEXT)

        for item in row["answers"]:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            head = table.rows[0].cells[0].merge(table.rows[0].cells[1])
            _shade(head, GREEN_LIGHT)
            head.text = ""
            p = head.paragraphs[0]
            r1 = p.add_run(_clean(
                f"Savol {label(item['section'], item['order_number'])}   ·   "
                f"{_fmt_answer_score(item)} ball"
            ))
            r1.font.bold = True
            r1.font.color.rgb = RGBColor.from_string(GREEN)
            head.add_paragraph(_clean(item["question"]))

            def add_row(label: str, value: str, color: str | None = None) -> None:
                cells = table.add_row().cells
                _cell_text(cells[0], label, bold=True, color=GRAY_TEXT, size=9)
                _cell_text(cells[1], value, color=color)

            add_row("Aytgani", item["transcript"] or "(transkripsiya qilinmadi)")
            if item["corrected_text"] and item["corrected_text"] != item["transcript"]:
                add_row("To'g'ri varianti", item["corrected_text"], color=GREEN)

            mistakes = [m for m in item["mistakes"] if isinstance(m, dict)]
            if mistakes:
                cells = table.add_row().cells
                _cell_text(cells[0], f"Xatolar ({len(mistakes)})", bold=True, color="B91C1C", size=9)
                cells[1].text = ""
                for i, m in enumerate(mistakes):
                    para = cells[1].paragraphs[0] if i == 0 else cells[1].add_paragraph()
                    wrong = para.add_run(_clean(m.get("original") or ""))
                    wrong.font.strike = True
                    wrong.font.color.rgb = RGBColor.from_string("B91C1C")
                    para.add_run("  →  ")
                    right = para.add_run(_clean(m.get("correct") or ""))
                    right.font.bold = True
                    right.font.color.rgb = RGBColor.from_string(GREEN)
                    explanation = m.get("explanation_uz") or m.get("explanation")
                    if explanation:
                        ex = para.add_run(_clean(f"  — {explanation}"))
                        ex.font.size = Pt(9)
                        ex.font.color.rgb = RGBColor.from_string(GRAY_TEXT)
            else:
                add_row("Xatolar", "Xato topilmadi", color=GREEN)

            _set_widths(table, [3.5, 23.2])
            doc.add_paragraph()


def build_report_docx(
    rows: list[dict[str, Any]],
    file_path: Path,
) -> None:
    """Natijalar hisobotini (.docx) yaratadi: statistika, natijalar jadvali
    (ball bo'yicha saralangan) va har bir foydalanuvchining javoblari."""
    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc = _setup_document()

    doc.add_heading("Turk tili speaking testi — natijalar", level=1)
    sub = doc.add_paragraph()
    srun = sub.add_run(f"Hisobot sanasi: {local_now():%Y-%m-%d %H:%M}")
    srun.font.size = Pt(10)
    srun.font.color.rgb = RGBColor.from_string(GRAY_TEXT)

    if not rows:
        doc.add_paragraph("Hech qanday test natijasi topilmadi.")
        doc.save(str(file_path))
        return

    rows = _sorted_rows(rows)
    _add_stats_cards(doc, _stats(rows))
    _add_summary_table(doc, rows)
    _add_answers_section(doc, rows)

    doc.save(str(file_path))
    logger.info("Hisobot fayli saqlandi: %s (%s qator)", file_path, len(rows))


# ───────────────────────────── Excel ─────────────────────────────

def build_report_xlsx(
    rows: list[dict[str, Any]],
    file_path: Path,
) -> None:
    """Natijalar (.xlsx): "Natijalar" — har kishi bir qatorda, har savol
    balli alohida ustunda; "Javoblar" — har javob bir qatorda."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    file_path.parent.mkdir(parents=True, exist_ok=True)
    rows = _sorted_rows(rows)

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=GREEN)
    zebra = PatternFill("solid", fgColor=ZEBRA)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, headers):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    def widths(ws, values):
        for i, w in enumerate(values, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()

    # ── Natijalar ──
    ws = wb.active
    ws.title = "Natijalar"
    qcols = _question_columns(rows)
    label = _labeler(rows)
    q_headers = [label(s, o) + (f"\n({mp} ball)" if mp else "") for s, o, mp in qcols]
    headers = ["#", "Ism familiya", "Telegram ID", "Telefon", "Username", "Holat",
               "Yakunlangan", f"Ball (/{MAX_SCORE})", "Daraja"] + q_headers
    style_header(ws, headers)

    for idx, row in enumerate(rows, start=1):
        by_q = {(a["section"], a["order_number"]): a for a in row["answers"]}
        q_values = []
        for s, o, mp in qcols:
            a = by_q.get((s, o))
            if a is None or a["score"] is None:
                q_values.append(None)
            else:
                q_values.append(a["earned"] if mp else a["score"])
        finished = to_local(row["finished_at"])
        ws.append([
            idx,
            _clean(row["full_name"]),
            row["telegram_id"],
            _clean(row["phone"]) if row["phone"] != "*" else "",
            f"@{_clean(row['username'])}" if row["username"] and row["username"] != "*" else "",
            _status_label(row["status"]),
            finished.replace(tzinfo=None) if finished else None,
            row["score"],
            row["level"] or "",
        ] + q_values)
        r = ws.max_row
        for cell in ws[r]:
            if idx % 2 == 0:
                cell.fill = zebra
        ws.cell(r, 7).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(r, 3).number_format = "0"
        for c in [1, 3, 7, 8, 9] + list(range(10, 10 + len(qcols))):
            ws.cell(r, c).alignment = center
        ws.cell(r, 8).font = Font(bold=True)
        colors = LEVEL_COLORS.get(row["level"] or "")
        if colors:
            ws.cell(r, 9).fill = PatternFill("solid", fgColor=colors[0])
            ws.cell(r, 9).font = Font(bold=True, color=colors[1])
    widths(ws, [5, 28, 14, 16, 18, 16, 17, 11, 11] + [9] * len(qcols))

    # ── Javoblar ──
    wa = wb.create_sheet("Javoblar")
    style_header(wa, ["Ism familiya", "Telegram ID", "Savol", "Savol matni", "Aytgani",
                      "To'g'ri varianti", "Xatolar", "Ball"])
    n = 0
    for row in rows:
        for a in row["answers"]:
            n += 1
            mistakes = [m for m in a["mistakes"] if isinstance(m, dict)]
            mistakes_text = "\n".join(
                f"{m.get('original') or ''} → {m.get('correct') or ''}"
                + (f" ({m.get('explanation_uz') or m.get('explanation')})"
                   if (m.get("explanation_uz") or m.get("explanation")) else "")
                for m in mistakes
            )
            wa.append([
                _clean(row["full_name"]),
                row["telegram_id"],
                label(a["section"], a["order_number"]),
                _clean(a["question"]),
                _clean(a["transcript"]),
                _clean(a["corrected_text"]) if a["corrected_text"] != a["transcript"] else "",
                _clean(mistakes_text),
                _fmt_answer_score(a),
            ])
            r = wa.max_row
            for cell in wa[r]:
                cell.alignment = wrap_top
                if n % 2 == 0:
                    cell.fill = zebra
            wa.cell(r, 2).number_format = "0"
    widths(wa, [26, 14, 8, 40, 50, 50, 45, 10])

    wb.save(str(file_path))
    logger.info("Excel hisobot saqlandi: %s", file_path)


# ───────────────────────────── Eksport ─────────────────────────────

def default_report_path(ext: str = "docx") -> Path:
    """Hozirgi vaqtga asoslangan hisobot fayli yo'li."""
    name = format_filename(f"natijalar_{local_now():%Y%m%d_%H%M%S}")
    return REPORTS_DIR / f"{name}.{ext}"


async def export_users_report(
    session: AsyncSession,
    file_path: Path | None = None,
) -> Path:
    """Barcha natijalarni .docx hisobotga eksport qiladi va yo'lini qaytaradi."""
    rows = await collect_users_with_results(session)
    target = file_path or default_report_path("docx")
    build_report_docx(rows, target)
    return target


async def export_users_reports(session: AsyncSession) -> list[Path]:
    """Word (.docx) va Excel (.xlsx) hisobotlarni yaratadi: [docx, xlsx]."""
    rows = await collect_users_with_results(session)
    docx_path = default_report_path("docx")
    xlsx_path = docx_path.with_suffix(".xlsx")
    build_report_docx(rows, docx_path)
    build_report_xlsx(rows, xlsx_path)
    return [docx_path, xlsx_path]
