"""
Kunlik Word hisobot testlari:
- Hisobotda foydalanuvchi javoblari (savol, transkript, xatolar, ball) bor
- Vaqtlar mahalliy (Toshkent) vaqtda ko'rsatiladi
- Keyingi mahalliy 00:00 gacha vaqt to'g'ri hisoblanadi
"""

import json
import os
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

import pytest
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.database.database import Base
from app.database.models import Answer, Question, TestAttempt, User
from app.services.user_report_exporter import export_users_report
from app.utils import helpers

TEST_DB_URL = os.getenv(
    "TEST_DB_URL",
    "postgresql+asyncpg://postgres:123@localhost:5432/turkish_bot_test",
)
test_engine = create_async_engine(TEST_DB_URL, echo=False, poolclass=NullPool)
TestSessionLocal = async_sessionmaker(test_engine, class_=AsyncSession, expire_on_commit=False)


@pytest.fixture(autouse=True)
async def setup_db():
    db_name = TEST_DB_URL.split("?", 1)[0].rsplit("/", 1)[-1]
    if not db_name.endswith("_test"):
        raise RuntimeError(f"Xavfsizlik: TEST_DB_URL test bazasi emas ({db_name})")

    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
        await conn.run_sync(Base.metadata.create_all)
    yield
    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


def _docx_text(path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.mark.anyio
async def test_report_contains_user_answers(tmp_path):
    async with TestSessionLocal() as session:
        user = User(
            telegram_id=777001, full_name="Ali Valiyev", username="ali",
            phone="+998901234567", is_registered=True,
        )
        idle = User(telegram_id=777002, full_name="Test Topshirmagan", is_registered=True)
        q1 = Question(section="1", order_number=1, text="Kendinizi tanıtın.", max_points=10)
        q2 = Question(section="1", order_number=2, text="Hobileriniz neler?")
        session.add_all([user, idle, q1, q2])
        await session.flush()

        attempt = TestAttempt(
            user_id=user.id, status="finished", score=52, level="B2",
            started_at=datetime(2026, 9, 25, 19, 0),  # UTC → 00:00 Toshkent
            finished_at=datetime(2026, 9, 25, 19, 10),
        )
        session.add(attempt)
        await session.flush()

        session.add_all([
            Answer(
                attempt_id=attempt.id, question_id=q2.id, score=60,
                transcript="Ben kitap okumak seviyorum.",
                feedback=json.dumps({"mistakes": []}),
            ),
            Answer(
                attempt_id=attempt.id, question_id=q1.id, score=80,
                transcript="Benim adım Ali. Ben öğrenciyim",
                feedback=json.dumps({
                    "corrected_text": "Benim adım Ali. Ben öğrenciyim.",
                    "mistakes": [{
                        "original": "öğrenciyim",
                        "correct": "öğrenciyim.",
                        "explanation_uz": "Nuqta tushib qolgan",
                    }],
                }, ensure_ascii=False),
            ),
        ])
        await session.commit()

    async with TestSessionLocal() as session:
        path = await export_users_report(session, tmp_path / "report.docx")

    text = _docx_text(path)

    # Umumiy jadval
    assert "Ali Valiyev" in text
    assert "Test Topshirmagan" in text
    assert "2026-09-26 00:10" in text  # yakunlangan vaqti Toshkent vaqtida
    assert "52/75" in text and "B2" in text

    # Javoblar bo'limi
    assert "Kendinizi tanıtın." in text
    assert "Benim adım Ali. Ben öğrenciyim" in text
    assert "Nuqta tushib qolgan" in text
    assert "8/10 ball" in text  # 80% × 10 ball
    assert "60% ball" in text   # max_points yo'q savol
    # Savol tartibi: 1.1 oldin, 1.2 keyin
    assert text.index("Kendinizi tanıtın.") < text.index("Hobileriniz neler?")
    # Test topshirmagan userning javoblar bo'limi yo'q
    assert "Test Topshirmagan —" not in text


@pytest.mark.anyio
async def test_empty_report(tmp_path):
    async with TestSessionLocal() as session:
        path = await export_users_report(session, tmp_path / "empty.docx")
    assert "Hech qanday test natijasi topilmadi." in _docx_text(path)


@pytest.mark.anyio
async def test_seconds_until_tashkent_midnight():
    tz = helpers.local_tz()
    assert datetime(2026, 1, 1, tzinfo=tz).utcoffset() == timedelta(hours=5)

    # 23:30 Toshkent → 30 daqiqa qoldi
    fake_now = datetime(2026, 9, 26, 23, 30, tzinfo=tz)
    with patch.object(helpers, "local_now", return_value=fake_now):
        assert helpers.seconds_until_local_midnight() == 30 * 60

    # 00:00 UTC = 05:00 Toshkent → 19 soat qoldi
    fake_now = datetime(2026, 9, 26, 0, 0, tzinfo=timezone.utc).astimezone(tz)
    with patch.object(helpers, "local_now", return_value=fake_now):
        assert helpers.seconds_until_local_midnight() == 19 * 3600


@pytest.mark.anyio
async def test_control_characters_do_not_break_report(tmp_path):
    async with TestSessionLocal() as session:
        user = User(telegram_id=777003, full_name="Ali\x0b Bad", is_registered=True)
        q = Question(section="1", order_number=1, text="Soru\x1f?")
        session.add_all([user, q])
        await session.flush()
        attempt = TestAttempt(user_id=user.id, status="finished", score=40, level="B1")
        session.add(attempt)
        await session.flush()
        session.add(Answer(
            attempt_id=attempt.id, question_id=q.id, score=50,
            transcript="Merhaba\x08 dünya",
            feedback=json.dumps({"mistakes": [{"original": "a\x01", "correct": "b"}]}),
        ))
        await session.commit()

    async with TestSessionLocal() as session:
        path = await export_users_report(session, tmp_path / "bad.docx")

    text = _docx_text(path)
    assert "Ali Bad" in text
    assert "Merhaba dünya" in text



@pytest.mark.anyio
async def test_excel_report_has_scores_and_answers(tmp_path):
    from openpyxl import load_workbook

    from app.services.user_report_exporter import build_report_xlsx, collect_users_with_results

    async with TestSessionLocal() as session:
        user = User(telegram_id=8963201482, full_name="Zarina", username="zar", is_registered=True)
        low = User(telegram_id=777005, full_name="Bekzod", is_registered=True)
        q1 = Question(section="1.1", order_number=1, text="Soru 1?", max_points=10)
        q2 = Question(section="1.2", order_number=1, text="Soru 2?")
        session.add_all([user, low, q1, q2])
        await session.flush()
        a1 = TestAttempt(user_id=user.id, status="finished", score=60, level="B2")
        a2 = TestAttempt(user_id=low.id, status="finished", score=30, level="Below B1")
        session.add_all([a1, a2])
        await session.flush()
        session.add_all([
            Answer(attempt_id=a1.id, question_id=q1.id, score=90, transcript="Merhaba",
                   feedback=json.dumps({"mistakes": [{"original": "x", "correct": "y"}]})),
            Answer(attempt_id=a1.id, question_id=q2.id, score=70, transcript="İyiyim"),
        ])
        await session.commit()

    async with TestSessionLocal() as session:
        rows = await collect_users_with_results(session)
    path = tmp_path / "r.xlsx"
    build_report_xlsx(rows, path)

    wb = load_workbook(path)
    ws = wb["Natijalar"]
    header = [c.value for c in ws[1]]
    assert header[:9] == ["#", "Ism familiya", "Telegram ID", "Telefon", "Username",
                          "Holat", "Yakunlangan", "Ball (/75)", "Daraja"]
    assert header[9:] == ["1.1\n(10 ball)", "1.2"]
    # Eng yuqori ball birinchi
    first = [c.value for c in ws[2]]
    assert first[1] == "Zarina" and first[2] == 8963201482 and first[7] == 60
    assert first[9:] == [9, 70]  # 90% × 10 ball = 9; ballsiz savol — foiz
    assert [c.value for c in ws[3]][1] == "Bekzod"

    wa = wb["Javoblar"]
    rows_a = [[c.value for c in r] for r in wa.iter_rows(min_row=2)]
    assert len(rows_a) == 2
    assert rows_a[0][2] == "1.1" and rows_a[0][6] == "x → y" and rows_a[0][7] == "9/10"
